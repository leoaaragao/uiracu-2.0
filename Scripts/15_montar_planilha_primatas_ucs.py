# -*- coding: utf-8 -*-
"""Monta a planilha final: Primatas x UCs, com abas por especie/UC/pais + metodologia."""
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

FONT = "Arial"
GREEN = "1F4E3D"
LIGHT = "E8F0EC"
GOLD = "B08D57"

header_font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color=GREEN, end_color=GREEN, fill_type="solid")
cell_font = Font(name=FONT, size=10)
title_font = Font(name=FONT, size=16, bold=True, color=GREEN)
sub_font = Font(name=FONT, size=10, italic=True, color="555555")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def escrever_df(ws, df, start_row=1, index=True, freeze=True):
    r = start_row
    headers = ([df.index.name or ""] if index else []) + list(df.columns.astype(str))
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=r, column=j, value=h)
        c.font = header_font
        c.fill = header_fill
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r += 1
    for idx, row in df.iterrows():
        vals = ([idx] if index else []) + list(row.values)
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = cell_font
            c.border = border
        r += 1
    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=(2 if index else 1))
    return r

wb = openpyxl.Workbook()

# ---------- Aba 1: Resumo / metodologia ----------
ws0 = wb.active
ws0.title = "Resumo e metodologia"
ws0.column_dimensions["A"].width = 34
ws0.column_dimensions["B"].width = 90
ws0["A1"] = "Primatas da Pan-Amazônia x Unidades de Conservação"
ws0["A1"].font = title_font
ws0.merge_cells("A1:B1")
ws0["A2"] = "Cruzamento espacial de ocorrências GBIF com as 85 UCs federais de estudo (AM/AC/RO) — Uiraçu 2.0"
ws0["A2"].font = sub_font
ws0.merge_cells("A2:B2")

linhas = [
    ("Data de geração", "24/09/2026"),
    ("Fonte dos dados", "GBIF.org — API pública occurrence/search (https://api.gbif.org/v1/occurrence/search)"),
    ("Espécies consultadas", "168 (todas as espécies aceitas em 18 gêneros de primatas amazônicos, obtidas do backbone do GBIF)"),
    ("Filtro geográfico", "Caixa delimitadora (WKT) ao redor das 85 UCs + margem de ~1 grau — SEM filtro de país (countryCode do GBIF pode estar incorreto; ver caso 'Loreto' no Diário de Bordo). A inclusão/exclusão real é decidida geometricamente contra o polígono de cada UC."),
    ("Estas ocorrências têm DOI citável?",
     "NÃO. Foram obtidas via API de busca (occurrence/search), que é síncrona e não gera DOI — diferente do download oficial usado para Lagothrix lagothricha (que tem DOI, via occurrence/download). Esta é uma varredura exploratória para identificar quais espécies têm mais incidência registrada nas UCs, não um dataset final citável. Se alguma espécie aqui identificada for usada de forma central em um relatório/artigo, gerar um download oficial específico para ela (mesmo processo já feito duas vezes para Lagothrix)."),
    ("Reprodutibilidade", "Scripts/12 (lista de espécies por gênero), Scripts/13 (download por espécie), Scripts/14 (cruzamento espacial com as UCs) — parâmetros exatos registrados em cada script."),
    ("Total de registros baixados (todas espécies, toda a caixa geográfica)", "14.588"),
    ("Espécies com pelo menos 1 registro na caixa geográfica", "123 de 168"),
    ("Registros que caem DENTRO do polígono de alguma das 85 UCs", "792"),
]
r = 4
for label, val in linhas:
    ws0.cell(row=r, column=1, value=label).font = Font(name=FONT, size=10, bold=True)
    c = ws0.cell(row=r, column=2, value=val)
    c.font = cell_font
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws0.row_dimensions[r].height = 60 if len(val) > 150 else (30 if len(val) > 60 else 18)
    r += 1

# ---------- Aba 2: Ranking por incidencia em UCs ----------
ranking = pd.read_csv("Referencias/ranking_especies_por_incidencia_ucs.csv").set_index("species")
ranking.columns = ["N de UCs distintas com registro", "N total de registros dentro das UCs"]
ranking.index.name = "Espécie"
ws1 = wb.create_sheet("Ranking por UC")
escrever_df(ws1, ranking)
ws1.column_dimensions["A"].width = 30
ws1.column_dimensions["B"].width = 22
ws1.column_dimensions["C"].width = 26

# ---------- Aba 3: Matriz especie x UC ----------
matriz = pd.read_csv("Referencias/matriz_especies_x_ucs_gbif.csv").set_index("species")
matriz.index.name = "Espécie"
ws2 = wb.create_sheet("Matriz Especie x UC")
escrever_df(ws2, matriz)
ws2.column_dimensions["A"].width = 30
for j in range(2, matriz.shape[1] + 2):
    ws2.column_dimensions[get_column_letter(j)].width = 14

# ---------- Aba 4: Especie x Pais (todos os registros baixados, nao so os dentro de UC) ----------
ocorrencias = pd.read_csv("Referencias/ocorrencias_primatas_brasil_gbif.csv")
tab_pais = pd.crosstab(ocorrencias["species"], ocorrencias["countryCode"])
tab_pais.index.name = "Espécie"
tab_pais["Total"] = tab_pais.sum(axis=1)
tab_pais = tab_pais.sort_values("Total", ascending=False)
ws3 = wb.create_sheet("Especie x Pais")
escrever_df(ws3, tab_pais)
ws3.column_dimensions["A"].width = 30
for j in range(2, tab_pais.shape[1] + 2):
    ws3.column_dimensions[get_column_letter(j)].width = 10

wb.save("Referencias/Primatas_x_UCs_GBIF.xlsx")
print("Salvo: Referencias/Primatas_x_UCs_GBIF.xlsx")
